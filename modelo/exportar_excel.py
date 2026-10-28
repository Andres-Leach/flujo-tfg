"""Vuelca el modelo a Excel (Modelo_financiero_Flujo.xlsx) para revisarlo sin Python."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import hipotesis as H

RAIZ = Path(__file__).resolve().parent.parent
R = json.loads((RAIZ / "modelo" / "resultados.json").read_text())
NEGRITA = Font(bold=True)
CABECERA = PatternFill("solid", fgColor="DCE6F2")
EUR = '#,##0 "€";-#,##0 "€"'
PCT = "0.0%"


def hoja(wb, titulo, cabecera, filas, formatos=None):
    ws = wb.create_sheet(titulo)
    ws.append(cabecera)
    for c in ws[1]:
        c.font, c.fill = NEGRITA, CABECERA
    for f in filas:
        ws.append(f)
        if isinstance(f[0], str) and f[0].isupper():
            for c in ws[ws.max_row]:
                c.font = NEGRITA
    for col in range(2, len(cabecera) + 1):
        for fila in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in fila:
                if isinstance(c.value, (int, float)):
                    c.number_format = (formatos or {}).get(ws.cell(c.row, 1).value, EUR)
    ws.column_dimensions["A"].width = 48
    for col in range(2, len(cabecera) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "B2"
    return ws


def pyg_filas(A):
    return [
        ["Ingresos por setup"] + A["ing_setup"],
        ["Ingresos por suscripción"] + A["ing_suscripcion"],
        ["Ingresos por plantillas"] + A["ing_plantillas"],
        ["Ingresos por cursos"] + A["ing_cursos"],
        ["Ingresos por la herramienta digital"] + A["ing_app"],
        ["IMPORTE NETO DE LA CIFRA DE NEGOCIOS"] + A["ingresos"],
        ["Costes variables (comisiones, alojamiento)"] + [-x for x in A["c_variables"]],
        ["MARGEN DE CONTRIBUCIÓN"] + A["margen_contribucion"],
        ["Cuota de autónomos (RETA)"] + [-x for x in A["reta"]],
        ["Otros gastos fijos de explotación"] + [-x for x in A["c_fijos"]],
        ["Marketing"] + [-x for x in A["marketing"]],
        ["Mantenimiento de la herramienta"] + [-x for x in A["mant_app"]],
        ["Gastos de puesta en marcha"] + [-x for x in A["puesta_marcha"]],
        ["EBITDA"] + A["ebitda"],
        ["Amortización del inmovilizado"] + [-x for x in A["amortizacion"]],
        ["RESULTADO DE EXPLOTACIÓN (EBIT)"] + A["ebit"],
        ["Gastos financieros"] + [-x for x in A["intereses"]],
        ["RESULTADO ANTES DE IMPUESTOS"] + A["rn_previo"],
        ["IRPF imputable a la actividad (estimado)"] + [-x for x in A["irpf"]],
        ["RESULTADO NETO DISPONIBLE PARA EL TITULAR"] + A["resultado_neto"],
    ]


def main():
    wb = Workbook()
    wb.remove(wb.active)
    hip = [
        ["Precio setup", H.PRECIO_SETUP], ["Precio suscripción mensual", H.PRECIO_SUSCRIPCION],
        ["Precio plantilla (neto)", H.PRECIO_PLANTILLA], ["Precio curso (neto)", H.PRECIO_CURSO],
        ["Precio herramienta (€/mes)", H.PRECIO_APP], ["Tasa de baja mensual", H.BAJA_MENSUAL],
        ["Horas por cliente y mes", H.HORAS_CLIENTE_MES], ["Horas por setup", H.HORAS_SETUP],
        ["Comisión de cobro (servicio)", H.COMISION_COBRO], ["Comisión plataforma (productos)", H.COMISION_PLATAFORMA],
        ["Aportación inicial del titular", H.APORTACION_INICIAL], ["Microcrédito (mes 25)", H.PRESTAMO],
        ["Tipo de interés del microcrédito", H.TIPO_PRESTAMO], ["Plazo del microcrédito (meses)", H.PLAZO_PRESTAMO],
        ["Tarifa plana RETA (€/mes)", H.TARIFA_PLANA], ["SMI anual 2026", H.SMI_ANUAL],
        ["Retención IRPF inicio de actividad", H.RETENCION], ["Tasa libre de riesgo", H.TASA_LIBRE_RIESGO],
        ["Prima de mercado", H.PRIMA_MERCADO], ["Beta", H.BETA], ["Prima específica", H.PRIMA_ESPECIFICA],
        ["Tasa de descuento", H.TASA_DESCUENTO], ["Crecimiento perpetuo (valor residual)", H.CRECIMIENTO_PERPETUO],
    ] + [[f"Coste fijo: {k} (€/mes)", v] for k, v in H.COSTES_FIJOS.items()]
    ws = hoja(wb, "Hipótesis", ["Hipótesis", "Valor"], hip)
    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in fila:
            nombre = ws.cell(c.row, 1).value.lower()
            c.number_format = PCT if any(p in nombre for p in ("tasa", "tipo", "comisión", "retención", "prima", "crecimiento")) else "#,##0.##"

    m = R["base"]["mensual"]
    claves = [("Altas de clientes", "altas"), ("Clientes activos", "activos"), ("Horas de servicio", "horas"),
              ("Usuarios de la herramienta", "usuarios_app"), ("Ingresos por setup", "ing_setup"),
              ("Ingresos por suscripción", "ing_suscripcion"), ("Ingresos por plantillas", "ing_plantillas"),
              ("Ingresos por cursos", "ing_cursos"), ("Ingresos por la herramienta", "ing_app"),
              ("INGRESOS TOTALES", "ingresos"), ("Costes variables", "c_variables"), ("Cuota RETA", "reta"),
              ("Otros gastos fijos", "c_fijos"), ("Marketing", "marketing"), ("Mantenimiento herramienta", "mant_app"),
              ("Amortización", "amortizacion"), ("Intereses", "intereses"), ("SALDO DE TESORERÍA", "caja")]
    fmt = {"Altas de clientes": "0.0", "Clientes activos": "0.0", "Horas de servicio": "0", "Usuarios de la herramienta": "0"}
    hoja(wb, "Mensual (base)", ["Concepto"] + [f"M{i}" for i in range(1, 37)],
         [[n] + m[k] for n, k in claves], fmt)

    for esc in ("base", "pesimista", "optimista"):
        A = R[esc]["anual"]
        hoja(wb, f"PyG {esc}", ["Cuenta de resultados", "Año 1", "Año 2", "Año 3"], pyg_filas(A))
        b = R[esc]["balance"]
        g = lambda k: [x[k] for x in b]
        hoja(wb, f"Balance {esc}", ["Balance a 31 de diciembre", "Año 1", "Año 2", "Año 3"], [
            ["Inmovilizado intangible (bruto)"] + g("intangible"), ["Inmovilizado material (bruto)"] + g("material"),
            ["Amortización acumulada"] + g("amort_acum"), ["ACTIVO NO CORRIENTE"] + g("inmov_neto"),
            ["Hacienda Pública deudora por IRPF"] + g("hp_deudora"), ["Tesorería"] + g("tesoreria"),
            ["TOTAL ACTIVO"] + g("activo"),
            ["Aportaciones del titular"] + g("aportaciones"), ["Resultados acumulados"] + g("resultados"),
            ["Retiradas del titular"] + g("retiradas"), ["PATRIMONIO NETO"] + g("pn"),
            ["Deudas con entidades de crédito a largo plazo"] + g("deuda_lp"),
            ["Deudas con entidades de crédito a corto plazo"] + g("deuda_cp"),
            ["Hacienda Pública acreedora por IVA"] + g("hp_iva"), ["Hacienda Pública acreedora por IRPF"] + g("hp_irpf"),
            ["TOTAL PASIVO"] + g("pasivo"),
            ["TOTAL PATRIMONIO NETO Y PASIVO"] + [x["pn"] + x["pasivo"] for x in b],
        ])
        F = A["flujo"]
        hoja(wb, f"Cash-flow {esc}", ["Estado de flujos de tesorería", "Año 1", "Año 2", "Año 3"], [
            ["Cobros de explotación (IVA incluido, netos de retenciones)"] + F["cobros_explotacion"],
            ["Pagos de explotación (IVA incluido)"] + [-x for x in F["pagos_explotacion"]],
            ["Liquidaciones de IVA"] + [-x for x in F["iva_neto"]],
            ["Pagos de IRPF (modelo 130 y renta)"] + [-x for x in F["irpf_pagado"]],
            ["FLUJO DE EXPLOTACIÓN"] + [F["cobros_explotacion"][a] - F["pagos_explotacion"][a] - F["iva_neto"][a]
                                         - F["irpf_pagado"][a] for a in range(3)],
            ["Inversiones (IVA incluido)"] + [-x for x in F["inversion"]],
            ["FLUJO DE INVERSIÓN"] + [-x for x in F["inversion"]],
            ["Aportación del titular"] + F["aportacion"], ["Microcrédito recibido"] + F["prestamo"],
            ["Devolución del principal"] + [-x for x in F["devolucion_prestamo"]],
            ["Retiradas del titular"] + [-x for x in F["retiradas"]],
            ["FLUJO DE FINANCIACIÓN"] + [F["aportacion"][a] + F["prestamo"][a] - F["devolucion_prestamo"][a]
                                          - F["retiradas"][a] for a in range(3)],
            ["TESORERÍA A FINAL DE AÑO"] + A["caja_fin"],
        ])

    filas = []
    for esc in ("pesimista", "base", "optimista"):
        v = R[esc]["valoracion"]
        filas += [[f"{esc.upper()}: flujos de caja libres (año 0 a 3)"] + v["flujos"],
                  [f"{esc}: flujos menos coste de oportunidad"] + v["flujos_oportunidad"]]
    hoja(wb, "Valoración", ["Concepto", "Año 0", "Año 1", "Año 2", "Año 3"], filas)
    ws = wb["Valoración"]
    ws.append([])
    ws.append(["Indicador", "Pesimista", "Base", "Optimista"])
    for c in ws[ws.max_row]:
        c.font, c.fill = NEGRITA, CABECERA
    for nombre, k, f in [("VAN (sin valor residual)", "van", EUR), ("VAN con valor residual", "van_con_vr", EUR),
                         ("TIR", "tir", PCT), ("Payback simple (años)", "payback", "0.00"),
                         ("Payback descontado (años)", "payback_desc", "0.00"),
                         ("VAN con coste de oportunidad", "van_oportunidad", EUR),
                         ("TIR con coste de oportunidad", "tir_oportunidad", PCT)]:
        ws.append([nombre] + [R[e]["valoracion"][k] for e in ("pesimista", "base", "optimista")])
        for c in ws[ws.max_row][1:]:
            c.number_format = f
    ws.append(["Tasa de descuento", R["tasa"]])
    ws.cell(ws.max_row, 2).number_format = PCT

    salida = RAIZ / "Modelo_financiero_Flujo.xlsx"
    wb.save(salida)
    print("guardado", salida)


if __name__ == "__main__":
    main()
